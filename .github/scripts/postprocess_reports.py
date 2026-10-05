from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path.cwd()
BASE = ROOT / 'Информатика и основы программирования' / 'Лабораторные работы'
HEADINGS_TO_KEEP = {
    'Назначение переменных',
    'Блок-схема алгоритма программы',
    'Текст программы',
    'Тестирование',
}


def _set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn('w:cantSplit')) is None:
        tr_pr.append(OxmlElement('w:cantSplit'))


def postprocess():
    """Keep program listings and their headings together across page breaks."""
    for lab in (3, 4, 5):
        path = BASE / f'ЛР{lab}' / 'report' / f'Лабораторная_№{lab}_вариант_8.docx'
        doc = Document(path)

        for paragraph in doc.paragraphs:
            if paragraph.text.strip() in HEADINGS_TO_KEEP:
                paragraph.paragraph_format.keep_with_next = True

        for table in doc.tables:
            text = '\n'.join(
                cell.text for row in table.rows for cell in row.cells
            )
            if '#include' not in text or 'int main' not in text:
                continue

            for row in table.rows:
                _set_cant_split(row)
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_together = True

        doc.save(path)


if __name__ == '__main__':
    postprocess()
