from pathlib import Path
import tempfile
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
ROOT=Path.cwd()
BASE=ROOT/'Информатика и основы программирования'/'Лабораторные работы'
TMP=Path(tempfile.mkdtemp(prefix='vogu_reports_'))
STUDENT='Маклаков Леонид Олегович'
GROUP='4Б09 РПС-11, 1 курс, 1 подгруппа'
TEACHER='Доцент Тихомирова Елена Николаевна'
DISCIPLINE='Информатика и основы программирования'
VARIANT=8
YEAR=2026

def set_cell_shading(cell, fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); trPr.append(x)
def set_run_font(run,name='Times New Roman',size=14,bold=None):
    run.font.name=name; run.font.size=Pt(size); run._element.rPr.rFonts.set(qn('w:eastAsia'),name)
    if bold is not None: run.bold=bold
def add_text(doc,text='',bold=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,first=1.25):
    p=doc.add_paragraph(); p.alignment=align; pf=p.paragraph_format; pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0); pf.first_line_indent=Cm(first) if first else None
    r=p.add_run(text); set_run_font(r,size=14,bold=bold); return p
def add_heading(doc,text,level=1):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; pf=p.paragraph_format; pf.space_before=Pt(8 if level==1 else 4); pf.space_after=Pt(3); pf.keep_with_next=True; pf.line_spacing=1.0
    r=p.add_run(text); set_run_font(r,size=14,bold=True); return p
def add_vars(doc,rows):
    t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; hdr=t.rows[0].cells
    for i,v in enumerate(['Переменная','Тип','Назначение']): hdr[i].text=v; set_cell_shading(hdr[i])
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=v
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for r in p.runs: set_run_font(r,size=11,bold=(ri==0))
    return t
def title_page(doc,lab,title):
    sec=doc.sections[0]; sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(3); sec.right_margin=Cm(1.5)
    for txt,bold in [('федеральное государственное бюджетное образовательное учреждение высшего образования',False),('«Вологодский государственный университет»',True),('',False),('Институт математики, естественных и компьютерных наук',False),('',False),('Кафедра автоматики и вычислительной техники',False)]:
        p=add_text(doc,txt,bold=bold,align=WD_ALIGN_PARAGRAPH.CENTER,first=0); p.paragraph_format.line_spacing=1.0
    for _ in range(2): doc.add_paragraph()
    p=add_text(doc,f'Отчет по лабораторной работе №{lab}',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,first=0); p.runs[0].font.size=Pt(16)
    add_text(doc,f'Дисциплина: «{DISCIPLINE}»',align=WD_ALIGN_PARAGRAPH.CENTER,first=0); add_text(doc,f'Название: «{title}»',align=WD_ALIGN_PARAGRAPH.CENTER,first=0); doc.add_paragraph()
    tab=doc.add_table(rows=1,cols=5); tab.style='Table Grid'; tab.alignment=WD_TABLE_ALIGNMENT.CENTER
    vals=['____09.03.04__.\nкод направления\nподготовки/специальности','__43.10__.\nкод выпускающей\nкафедры','______1______\nрегистрационный номер по журналу','_____1_____\nкод формы\nобучения',f'___{YEAR}___\nгод']
    for c,v in zip(tab.rows[0].cells,vals):
        c.text=v
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: set_run_font(r,size=9)
    doc.add_paragraph(); info=doc.add_table(rows=5,cols=2); info.alignment=WD_TABLE_ALIGNMENT.CENTER
    vals=[('Руководитель',TEACHER),('Выполнил студент',STUDENT),('Группа, курс',GROUP),('Оценка','____________________'),('Дата','____________________')]
    for i,(a,b) in enumerate(vals): info.cell(i,0).text=a; info.cell(i,1).text=b
    for row in info.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs: set_run_font(r,size=12)
    for _ in range(3): doc.add_paragraph()
    add_text(doc,'Вологда',align=WD_ALIGN_PARAGRAPH.CENTER,first=0); add_text(doc,f'{YEAR} г.',align=WD_ALIGN_PARAGRAPH.CENTER,first=0); doc.add_page_break()
def font_mono(size=18):
    for p in ['/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf','/usr/share/fonts/truetype/liberation2/LiberationMono-Regular.ttf']:
        if Path(p).exists(): return ImageFont.truetype(p,size)
    return ImageFont.load_default()
def font_sans(size=20):
    p='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    return ImageFont.truetype(p,size) if Path(p).exists() else ImageFont.load_default()
def wrap(draw,text,font,maxw):
    words=text.split(); lines=[]; cur=''
    for w in words:
        cand=(cur+' '+w).strip()
        if draw.textbbox((0,0),cand,font=font)[2]<=maxw: cur=cand
        else:
            if cur: lines.append(cur)
            cur=w
    if cur: lines.append(cur)
    return lines or ['']
def flow_png(nodes,path):
    W=1200; boxw=850; y=50; font=font_sans(24); temp=Image.new('RGB',(W,3000),'white'); d=ImageDraw.Draw(temp); pos=[]
    for kind,text in nodes:
        lines=wrap(d,text,font,boxw-100); h=max(90,40+34*len(lines)); pos.append((kind,y,h,lines)); y+=h+80
    img=Image.new('RGB',(W,y+20),'white'); d=ImageDraw.Draw(img)
    for idx,(kind,y,h,lines) in enumerate(pos):
        x=(W-boxw)//2; x2=x+boxw
        if kind in ('start','end'): d.rounded_rectangle((x,y,x2,y+h),radius=h//2,outline='black',width=3)
        elif kind=='dec':
            cx=W//2; d.polygon([(cx,y),(x2,y+h//2),(cx,y+h),(x,y+h//2)],outline='black')
        elif kind=='io':
            off=60; d.polygon([(x+off,y),(x2,y),(x2-off,y+h),(x,y+h)],outline='black')
        else: d.rectangle((x,y,x2,y+h),outline='black',width=3)
        yy=y+(h-34*len(lines))//2
        for line in lines:
            bb=d.textbbox((0,0),line,font=font); d.text(((W-(bb[2]-bb[0]))//2,yy),line,font=font,fill='black'); yy+=34
        if idx<len(pos)-1:
            y2=pos[idx+1][1]; cx=W//2; d.line((cx,y+h,cx,y2-12),fill='black',width=3); d.polygon([(cx,y2),(cx-8,y2-14),(cx+8,y2-14)],fill='black')
    img.save(path,optimize=True)
def terminal_png(text,path):
    text=text.replace('\t','    ').strip('\n'); lines=text.splitlines() or ['(нет вывода)']; font=font_mono(20); margin=28; lineh=30
    W=max(900,min(1500,max((len(x) for x in lines),default=20)*13+2*margin)); H=max(220,len(lines)*lineh+2*margin); im=Image.new('RGB',(W,H),(25,25,25)); d=ImageDraw.Draw(im); y=margin
    for line in lines: d.text((margin,y),line,font=font,fill=(240,240,240)); y+=lineh
    im.save(path,optimize=True)
def add_picture_center(doc,path,width=15.5):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3); p.add_run().add_picture(str(path),width=Cm(width))
def add_code(doc,code):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; p=t.cell(0,0).paragraphs[0]; p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
    for line in code.rstrip().splitlines(): set_run_font(p.add_run(line+'\n'),name='Courier New',size=9)
def build_lab(lab,spec):
    doc=Document(); title_page(doc,lab,spec['title']); add_heading(doc,f'Вариант №{VARIANT}'); add_heading(doc,'Цель работы'); add_text(doc,spec['goal'])
    for idx,task in enumerate(spec['tasks'],1):
        add_heading(doc,f'Задание №{idx}'); add_heading(doc,'Текст задания',2); add_text(doc,task['text'])
        if task.get('note'): add_text(doc,task['note'])
        add_heading(doc,'Работа программы',2); add_text(doc,task['desc']); add_heading(doc,'Назначение переменных',2); add_vars(doc,task['vars']); add_heading(doc,'Блок-схема алгоритма программы',2)
        fp=TMP/f'flow_{lab}_{idx}.png'; flow_png(task['flow'],fp); add_picture_center(doc,fp,14.5); add_heading(doc,'Текст программы',2)
        src=BASE/f'ЛР{lab}'/'src'/f'task{idx}.cpp'; add_code(doc,src.read_text(encoding='utf-8')); add_heading(doc,'Тестирование',2)
        for j,out in enumerate(task['outputs'],1):
            screen=TMP/f'test_{lab}_{idx}_{j}.png'; terminal_png(out,screen); add_picture_center(doc,screen,15.5); cap=add_text(doc,f'Рисунок {idx}.{j} – Результат тестирования программы',align=WD_ALIGN_PARAGRAPH.CENTER,first=0); cap.paragraph_format.line_spacing=1.0
    add_heading(doc,'Вывод'); add_text(doc,spec['conclusion']); outdir=BASE/f'ЛР{lab}'/'report'; outdir.mkdir(parents=True,exist_ok=True); out=outdir/f'Лабораторная_№{lab}_вариант_8.docx'; doc.save(out); print(out)
